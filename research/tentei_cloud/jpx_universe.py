#!/usr/bin/env python3
"""Snapshot the current JPX domestic common-share universe for Cloud research."""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import urllib.parse
import urllib.request
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import load_workbook

JPX_PAGE = "https://www.jpx.co.jp/markets/statistics-equities/misc/01.html"
UA = "Mozilla/5.0 TenteiCloudResearch/1.0"
MARKETS = {
    "プライム（内国株式）": "Prime",
    "スタンダード（内国株式）": "Standard",
    "グロース（内国株式）": "Growth",
}
CODE = re.compile(r"^[0-9A-Z]{4}$")


class ExcelLinks(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag != "a":
            return
        href = dict(attrs).get("href", "")
        if re.search(r"\.(?:xlsx|xls)(?:\?|$)", href, re.I):
            self.links.append(href)


def discover_excel(page_html: str, page_url: str = JPX_PAGE) -> str:
    parser = ExcelLinks()
    parser.feed(page_html)
    candidates = [urllib.parse.urljoin(page_url, h) for h in parser.links]
    candidates = [u for u in candidates if urllib.parse.urlparse(u).hostname == "www.jpx.co.jp"]
    selected = [u for u in candidates if re.search(r"/data_j\.xlsx?(?:\?|$)", u, re.I)]
    if len(selected) != 1:
        raise ValueError(f"expected one official data_j Excel link; found {len(selected)}")
    return selected[0]


def extract_universe(excel_bytes: bytes) -> tuple[list[dict[str, str]], str]:
    book = load_workbook(io.BytesIO(excel_bytes), read_only=True, data_only=True)
    sheet = book.active
    rows = sheet.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows)]
    for required in ("日付", "コード", "銘柄名", "市場・商品区分"):
        if required not in header:
            raise ValueError(f"JPX column missing: {required}")
    idx = {name: header.index(name) for name in ("日付", "コード", "銘柄名", "市場・商品区分")}
    out: dict[str, dict[str, str]] = {}
    dates = set()
    for row in rows:
        segment = str(row[idx["市場・商品区分"]] or "").strip()
        market = MARKETS.get(segment)
        if market is None:
            continue
        code = str(row[idx["コード"]] or "").strip().upper()
        name = str(row[idx["銘柄名"]] or "").strip()
        # JPX groups listed preferred/class shares under domestic stocks too.
        if "優先株式" in name or "種類株式" in name:
            continue
        if not CODE.fullmatch(code):
            raise ValueError(f"invalid JPX ordinary-stock code: {code!r}")
        entry = {"code": code, "name": name,
                 "market": market, "ticker": code + ".T"}
        if code in out and out[code] != entry:
            raise ValueError(f"conflicting duplicate JPX code: {code}")
        out[code] = entry
        dates.add(str(row[idx["日付"]]).split(".")[0])
    if not out or len(dates) != 1:
        raise ValueError("JPX universe empty or mixed snapshot dates")
    return [out[k] for k in sorted(out)], dates.pop()


def fetch(url: str, referer: str = "") -> bytes:
    headers = {"User-Agent": UA, "Accept-Language": "ja,en-US;q=0.8"}
    if referer:
        headers["Referer"] = referer
    with urllib.request.urlopen(urllib.request.Request(url, headers=headers), timeout=60) as response:
        return response.read()


def write_snapshot(rows: list[dict[str, str]], asof: str, source_url: str,
                   source_bytes: bytes, outdir: Path) -> dict:
    outdir.mkdir(parents=True, exist_ok=True)
    with (outdir / "jpx_universe.csv").open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["code", "name", "market", "ticker"])
        writer.writeheader()
        writer.writerows(rows)
    (outdir / "symbols.txt").write_text("".join(r["code"] + "\n" for r in rows), encoding="utf-8")
    counts = {m: sum(r["market"] == m for r in rows) for m in ("Prime", "Standard", "Growth")}
    meta = {"source_page": JPX_PAGE, "source_excel": source_url,
            "source_sha256": hashlib.sha256(source_bytes).hexdigest(),
            "snapshot_date": asof, "target_symbols": len(rows), "market_counts": counts}
    (outdir / "jpx_universe_meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    return meta


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--outdir", required=True)
    a = ap.parse_args()
    page = fetch(JPX_PAGE).decode("utf-8")
    excel_url = discover_excel(page)
    excel = fetch(excel_url, JPX_PAGE)
    rows, asof = extract_universe(excel)
    print(json.dumps(write_snapshot(rows, asof, excel_url, excel, Path(a.outdir)), ensure_ascii=False))


if __name__ == "__main__":
    main()
